"""
Document Processor for BIS SmartSpec AI
Handles PDF, DOCX, and TXT file parsing.
Extracts text and detects IS standard references.
"""
import os
import re
import uuid
from typing import Tuple, List
from loguru import logger


class DocumentProcessor:
    """Extracts text from uploaded documents."""

    MAX_TEXT_LENGTH = 100_000  # characters

    def process(self, file_path: str, file_type: str) -> Tuple[str, List[dict]]:
        """
        Extract text and IS references from document.
        
        Returns:
            (extracted_text, detected_standards)
        """
        text = ""
        
        try:
            if file_type == "pdf":
                text = self._extract_pdf(file_path)
            elif file_type == "docx":
                text = self._extract_docx(file_path)
            elif file_type == "txt":
                text = self._extract_txt(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")
        except Exception as e:
            logger.error(f"Document extraction failed: {e}")
            raise

        # Truncate if too long
        if len(text) > self.MAX_TEXT_LENGTH:
            text = text[:self.MAX_TEXT_LENGTH] + "\n[Document truncated for processing]"

        # Extract IS references
        from app.utils.language import extract_is_references
        detected = extract_is_references(text)

        return text.strip(), detected

    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF using pypdf."""
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            pages = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    pages.append(page_text)
            return "\n".join(pages)
        except ImportError:
            raise ImportError("pypdf not installed. Cannot process PDF files.")

    def _extract_docx(self, file_path: str) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text.strip())
            return "\n".join(paragraphs)
        except ImportError:
            raise ImportError("python-docx not installed. Cannot process DOCX files.")

    def _extract_txt(self, file_path: str) -> str:
        """Extract text from TXT file."""
        encodings = ["utf-8", "utf-16", "latin-1", "cp1252"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        with open(file_path, "rb") as f:
            return f.read().decode("utf-8", errors="ignore")

    def generate_summary(self, text: str, detected_standards: List[dict]) -> dict:
        """Generate a document analysis summary."""
        # Count words
        word_count = len(text.split())
        
        # Extract potential product/specification keywords
        product_indicators = self._extract_product_terms(text)
        
        return {
            "word_count": word_count,
            "detected_standards_count": len(detected_standards),
            "detected_standards": detected_standards,
            "product_indicators": product_indicators[:10],
            "text_preview": text[:500] + "..." if len(text) > 500 else text,
        }

    def _extract_product_terms(self, text: str) -> List[str]:
        """Extract potential product/specification terms."""
        # Common procurement keywords
        patterns = [
            r'\b(?:supply of|procurement of|purchase of|supply and installation of)\s+([^,.\n]{5,50})',
            r'\b(?:specification for|specifications for)\s+([^,.\n]{5,50})',
            r'\b(?:material|equipment|product|item|goods):\s*([^,.\n]{5,50})',
        ]
        
        terms = []
        for pattern in patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            terms.extend([m.strip() for m in matches])
        
        return list(set(terms))[:10]


# Singleton
_processor: DocumentProcessor = None


def get_document_processor() -> DocumentProcessor:
    global _processor
    if _processor is None:
        _processor = DocumentProcessor()
    return _processor
